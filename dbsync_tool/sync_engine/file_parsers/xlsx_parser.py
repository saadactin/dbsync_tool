"""XLSX parser using openpyxl read-only streaming."""
from __future__ import annotations

from typing import Dict, Iterable, List, Tuple

from sync_engine.exceptions import FlatFileReadError

from .common import normalize_headers, resolve_source_path


def _import_openpyxl():
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise FlatFileReadError(
            "openpyxl is required for XLSX sources. Install requirements.txt."
        ) from exc
    return load_workbook


def _stringify(value) -> str:
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'true' if value else 'false'
    return str(value)


def iter_records(
    file_source,
    *,
    chunk_size: int = 1000,
    encoding: str = 'utf-8',  # ignored for XLSX; kept for contract parity
) -> Tuple[List[str], Iterable[List[Dict[str, str]]], Dict[str, int]]:
    del encoding  # XLSX is binary; encoding does not apply.
    if chunk_size <= 0:
        chunk_size = 1000
    path = resolve_source_path(file_source)
    sheet_name = (getattr(file_source, 'sheet_name', '') or '').strip()
    has_header = bool(getattr(file_source, 'has_header', True))

    stats: Dict[str, int] = {
        'rows_read': 0,
        'rows_padded': 0,
        'rows_truncated': 0,
        'single_column_warning': 0,
    }

    load_workbook = _import_openpyxl()

    try:
        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise FlatFileReadError(f"Failed to open XLSX file: {exc}") from exc

    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise FlatFileReadError(
                    f"Sheet '{sheet_name}' not found. Available: {workbook.sheetnames}"
                )
            sheet = workbook[sheet_name]
        else:
            sheet = workbook[workbook.sheetnames[0]]

        rows_iter = sheet.iter_rows(values_only=True)
        try:
            first = next(rows_iter)
        except StopIteration:
            first = None

        if first is None:
            headers: List[str] = []
        elif has_header:
            headers = normalize_headers([_stringify(c) for c in first])
        else:
            headers = [f'column_{idx + 1}' for idx in range(len(first))]
    finally:
        # We deliberately keep workbook open for streaming; close in iterator.
        pass

    if len(headers) == 1:
        stats['single_column_warning'] = 1

    initial_first_row_data = first if (first is not None and not has_header) else None

    def _row_to_dict(raw) -> Dict[str, str]:
        if raw is None:
            raw = ()
        row_vals = [_stringify(c) for c in raw]
        if len(row_vals) < len(headers):
            row_vals.extend([''] * (len(headers) - len(row_vals)))
            stats['rows_padded'] += 1
        if len(row_vals) > len(headers):
            row_vals = row_vals[: len(headers)]
            stats['rows_truncated'] += 1
        stats['rows_read'] += 1
        return {headers[idx]: row_vals[idx] for idx in range(len(headers))}

    def _iterator():
        try:
            chunk: List[Dict[str, str]] = []
            if initial_first_row_data is not None:
                chunk.append(_row_to_dict(initial_first_row_data))
            for raw in rows_iter:
                chunk.append(_row_to_dict(raw))
                if len(chunk) >= chunk_size:
                    yield chunk
                    chunk = []
            if chunk:
                yield chunk
        finally:
            try:
                workbook.close()
            except Exception:
                pass

    return headers, _iterator(), stats
